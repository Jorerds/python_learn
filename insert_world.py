# 将应word模板的占位符名字对应的图片插入到word里面，并且生成一个新的word文档
import re,os,glob,datetime
from docx import Document
from docx.shared import Inches,Cm


def docx_work(mod_path,img_list,save_path):
    """
    读取模板wokd文件
    :param mod_path:模板文件的路径
    :param img_list:图片文件列表
    :param save_path:文档保存目录
    :return:
    """
    file_path=mod_path
    doc=Document(file_path)

    for para in doc.paragraphs:
        if img_list:
            # 循环图片文件列表
            for img_file_path in img_list:
                # 将文件完整的路径去掉，只保留文件名和后缀
                file_name_str = os.path.basename(img_file_path)
                # 将文件名和后缀切分成集合，下标0则是文件名
                file_name = os.path.splitext(file_name_str)[0]
                placeholder='{{%s}}'%file_name
                # 找到模板中对应的占位符
                if placeholder in para.text:
                    print(f"正在为模板替换{file_name}图表")
                    # 替换文本，防止占位符残留
                    para.text=para.text.replace(placeholder,"")
                    # 在找到的占位符对象段落中添加新的对象来放图片
                    # 这样图片会紧跟段落文字后面
                    new_run=para.add_run()
                    # 设置图片的大小尺寸
                    new_run.add_picture(img_file_path,width=Cm(14))
    # 获取当前日期时间
    new_date_time=datetime.datetime.now()
    # 年份
    new_year = new_date_time.year
    # 月份
    new_month = new_date_time.month
    curr_file=save_path
    file_save=os.path.join(curr_file,f'{new_year}年{new_month}月份发电机月报.docx')
    # 保存为新的文件
    doc.save(file_save)

    return "文档生成成功！"